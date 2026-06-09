from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from typer.testing import CliRunner

from app import cli
from config import OUTPUT_SHEETS, SCHEMA_GUIDE_SHEET
from generator.pipeline import generate_star_schema


def test_export_workbook_contains_expected_sheets(
    tmp_path: Path, sample_raw_frame: pd.DataFrame
) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "star.xlsx"
    sample_raw_frame.to_excel(input_path, sheet_name="Raw Data", index=False)

    result = generate_star_schema(input_path, output_path)
    workbook = load_workbook(result.output_path, read_only=True)

    assert workbook.sheetnames == OUTPUT_SHEETS
    assert workbook.sheetnames[0] == SCHEMA_GUIDE_SHEET
    assert workbook["FACT_Attendance"].max_row == 3
    assert workbook["FACT_Attendance"]["A1"].value == "Employee_ID"


def test_cli_smoke_generates_workbook(
    tmp_path: Path, sample_raw_frame: pd.DataFrame
) -> None:
    input_path = tmp_path / "input.xlsx"
    output_path = tmp_path / "cli-output.xlsx"
    sample_raw_frame.to_excel(input_path, sheet_name="Raw Data", index=False)

    runner = CliRunner()
    result = runner.invoke(cli, [str(input_path), "--output", str(output_path)])

    assert result.exit_code == 0
    assert output_path.exists()
    assert "generated successfully" in result.output
